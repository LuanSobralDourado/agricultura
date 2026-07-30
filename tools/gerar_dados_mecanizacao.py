# -*- coding: utf-8 -*-
"""
Gera js/dados-mecanizacao.js a partir da aba "dados" de mecanizacao26.xlsx.

Uso (na raiz do projeto):
    python tools/gerar_dados_mecanizacao.py

Somente a aba "dados" e considerada. Campos sensiveis (CPF, telefone,
data de nascimento, e-mail) NAO sao exportados: o CPF e usado apenas para
gerar um id anonimo de produtor, que permite contar produtores distintos.
"""
import io
import json
import os
import re
import unicodedata
from datetime import datetime

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(RAIZ, "mecanizacao26.xlsx")
SAIDA = os.path.join(RAIZ, "js", "dados-mecanizacao.js")
ABA = "dados"

NAO_INF = "Não informado"


# ----------------------------------------------------------------- utilidades
def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def eh_nao_informado(s):
    t = sem_acento(s).lower().strip(' "\'.')
    return t in ("", "nao informado", "nao informada", "nao informando", "nao inofrmado",
                 "nao", "n/a", "na", "0", "-")


def rotulo(s):
    """Normaliza um rotulo de categoria: colapsa variantes de 'nao informado'."""
    s = texto(s)
    return NAO_INF if eh_nao_informado(s) else s


def titulo(s):
    return " ".join(p.capitalize() if len(p) > 2 else p.lower() for p in s.split())


def numero(v):
    """Converte para float, ignorando lixo ('00000', '0000000000', texto)."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = texto(v).replace(".", "").replace(",", ".") if texto(v).count(",") else texto(v)
    if not re.fullmatch(r"-?\d+(\.\d+)?", t or ""):
        return 0.0
    n = float(t)
    # sequencias de zeros usadas como preenchimento
    return 0.0 if re.fullmatch(r"0+(\.0+)?", t) else n


def pid_cpf(cpf):
    """Id anonimo e ESTAVEL do produtor, derivado do CPF (FNV-1a de 32 bits,
    duas passadas combinadas em 53 bits — cabe num Number do JavaScript).

    Precisa ser estavel porque os dados de 2026 e os anos anteriores sao
    gerados por scripts diferentes: contar produtores distintos no consolidado
    so funciona se o mesmo CPF virar o mesmo id nos dois arquivos.
    O mesmo algoritmo esta em js/importar.js — mexeu aqui, mexa la."""
    def fnv(base):
        h = base
        for ch in cpf:
            h = ((h ^ ord(ch)) * 16777619) & 0xFFFFFFFF
        return h
    return fnv(2166136261) * 2097152 + (fnv(2166136269) & 0x1FFFFF)


def data_iso(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    t = texto(v)
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", t)
    if m:
        return "%04d-%02d-%02d" % (int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return ""


# ------------------------------------------------------ normalizacoes de dominio
def norm_escritorio(s):
    s = texto(s)
    s = re.sub(r"^Escrit[óo]rio\s+Local\s+(de|da|do)?\s*", "", s, flags=re.I)
    chave = sem_acento(s).lower()
    mapa = {
        "transacrena": "Transacreana",
        "transacreana": "Transacreana",
        "brasileia": "Brasiléia",
    }
    return mapa.get(chave, titulo(s) if s.isupper() else s) or NAO_INF


def norm_estado_civil(s):
    c = sem_acento(texto(s)).lower()
    if c.startswith("casado"):
        return "Casado(a)"
    if c.startswith("solt") or c.startswith("solit"):
        return "Solteiro(a)"
    if c.startswith("amasiado"):
        return "Amasiado(a)"
    if c.startswith("viuvo"):
        return "Viúvo(a)"
    if c.startswith("divorciado"):
        return "Divorciado(a)"
    if c.startswith("separado"):
        return "Separado(a)"
    return NAO_INF


def norm_tecnico(s):
    s = texto(s)
    if s.isupper():
        s = titulo(s)
    c = sem_acento(s).lower()
    mapa = {
        "jorge ney pontes": "Jorge Ney Pontes Araújo",
        "jorge ney pontes araujo": "Jorge Ney Pontes Araújo",
        "antonio francisco": "Antônio Francisco de Araújo do Nascimento",
    }
    return mapa.get(c, s) or NAO_INF


# Extracao por palavra-chave: os campos livres misturam varios valores
# numa mesma celula ("Grade aradora Grade niveladora Plantadeira").
IMPLEMENTOS = [
    ("Grade aradora", ("grade aradora", "grade de arado", "aradora", "arado")),
    ("Grade niveladora", ("niveladora", "nivelado")),
    ("Plantadeira", ("plantadeira",)),
    ("Colheitadeira", ("colheitadeira",)),
    ("Pulverizador", ("pulverizador",)),
    ("Jogadora de calcário", ("calcario",)),
    ("Destoca", ("destoca",)),
    ("Roçadeira", ("rocad",)),
    ("Piscicultura", ("piscicultura", "psicultura", "pisicultura", "piscicultuta", "picicultura")),
    ("Tanque / açude", ("tanque", "acude", "reforma")),
]

MAQUINAS = [
    ("Escavadeira hidráulica", ("escavadeira",)),
    ("Pá carregadeira", ("pa carregadeira", "pa mecanica", "pa carregaderira", "pa mecanic")),
    ("Trator de esteira", ("esteira",)),
    ("Trator de pneu", ("pneu",)),
    ("Trator agrícola", ("trator agricola",)),
    ("Retroescavadeira", ("retroescavadeira",)),
    ("New Holland", ("new holland",)),
    ("Massey Ferguson", ("massey",)),
    ("John Deere", ("john deere",)),
    ("Solis 90", ("solis", "soleis", "soles")),
]


def extrair(valor, tabela):
    """Devolve a lista de categorias canonicas presentes no texto livre."""
    c = sem_acento(texto(valor)).lower()
    if not c or eh_nao_informado(texto(valor)):
        return []
    achados = []
    for canonico, chaves in tabela:
        if any(k in c for k in chaves) and canonico not in achados:
            achados.append(canonico)
    return achados


# ------------------------------------------------------------------- leitura
def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    if ABA not in wb.sheetnames:
        raise SystemExit('Aba "%s" nao encontrada em %s' % (ABA, XLSX))
    ws = wb[ABA]

    linhas = list(ws.iter_rows(values_only=True))
    cab = [texto(h) for h in linhas[0]]
    dados = [r for r in linhas[1:] if any(c is not None and texto(c) != "" for c in r)]

    def idx(nome, ocorrencia=0):
        achados = [i for i, h in enumerate(cab) if h == nome]
        return achados[ocorrencia]

    I = {
        "carimbo": idx("Carimbo de data/hora"),
        "escritorio": idx("Escritório Local"),
        "vistoria": idx("Data da Vistoria"),
        "tecnico": idx("Nome do responsável técnico"),
        "produtor": idx("Nome do produtor"),
        "sexo": idx("Sexo"),
        "cpf": idx("CPF"),
        "civil": idx("Estado Civil"),
        "assoc": idx("Nome da Associação/Cooperativa:"),
        "dap": idx("Possui DAP:"),
        "municipio": idx("Município:"),
        "endereco": idx("Endereço (Projeto de Assentamento/Comunidade, BR/Ramal, Km, nº do lote):"),
        "propriedade": idx("Nome da Propriedade:"),
        "ponto": idx("Ponto de controle"),
        "horas": idx("Quantidade de horas"),
        "maquina": idx("Máquina"),
        "total_mec": idx("Total mecanizado"),
        "trator": idx("Nome do trator"),
        "tipo_trator": idx("Tipo de trator"),
        "tipo_impl": idx("Tipo de Implemento"),
        "nome_impl": idx("Nome do implemento"),
        "acudes": idx("Quantidade de açudes"),
        "formulario": idx("Formulário de Mecanização"),
        "obs": idx("Informe a observação"),
    }
    CULTURAS = [(idx("Primeira cultura"), idx("Área (hectare)", 0), idx("Sistema de Cultivo ".strip(), 0)),
                (idx("Segunda cultura"), idx("Área (hectare)", 1), idx("Sistema de Cultivo", 1)),
                (idx("Terceira cultura"), idx("Área (hectare)", 2), idx("Sistema de Cultivo", 2)),
                (idx("Quarta cultura"), idx("Área (hectare)", 3), idx("Sistema de Cultivo", 3))]
    DAE_VALORES = [idx("Informe o valor da DAE - %d" % n) for n in range(1, 11)]

    cpfs = {}
    registros = []
    qualidade = {"cpf_invalido": 0, "sem_data_valida": 0, "vistoria_outro_ano": 0,
                 "sem_geo": 0, "sem_formulario": 0, "acudes_texto": 0}

    i_geox, i_geoy = idx("Geo X"), idx("Geo Y")

    for r in dados:
        def v(k):
            j = I[k]
            return r[j] if j < len(r) else None

        cpf = texto(v("cpf"))
        if not re.fullmatch(r"\d{3}\.\d{3}\.\d{3}-\d{2}", cpf) or cpf == "***.***.***-**":
            qualidade["cpf_invalido"] += 1
            pid = -1
        else:
            pid = pid_cpf(cpf)
            cpfs[cpf] = pid

        # A data de referencia do painel e a de INSERCAO (Carimbo de data/hora):
        # e a unica confiavel na planilha. A Data da Vistoria vai junto, apenas
        # para exibicao, porque tem digitacao errada e vistorias de anos anteriores.
        data = data_iso(v("carimbo"))
        vistoria = data_iso(v("vistoria"))
        if not data:
            data = vistoria
            qualidade["sem_data_valida"] += 1
        if vistoria and vistoria[:4] != data[:4]:
            qualidade["vistoria_outro_ano"] += 1

        culturas = []
        for ic, ia, isis in CULTURAS:
            nome = texto(r[ic]) if ic < len(r) else ""
            if not nome or eh_nao_informado(nome):
                continue
            culturas.append([rotulo(nome), round(numero(r[ia]), 2), rotulo(r[isis])])

        dae = round(sum(numero(r[j]) for j in DAE_VALORES if j < len(r)), 2)

        gx, gy = numero(r[i_geox] if i_geox < len(r) else None), numero(r[i_geoy] if i_geoy < len(r) else None)
        if not gx or not gy:
            qualidade["sem_geo"] += 1

        form = texto(v("formulario"))
        if not form.startswith("http"):
            form = ""
            qualidade["sem_formulario"] += 1

        # "Quantidade de acudes" as vezes vem como texto ('01', '05'). O painel
        # converte e conta; o SOMA do Excel ignora essas celulas — dai a
        # diferenca entre os dois totais. Contamos para poder avisar.
        bruto_ac = v("acudes")
        if not isinstance(bruto_ac, (int, float)) and numero(bruto_ac) > 0:
            qualidade["acudes_texto"] += 1

        registros.append({
            "d": data,        # data de insercao (Carimbo de data/hora)
            "ex": data[:4],   # exercicio (ano) a que o registro pertence
            "dv": vistoria,   # data da vistoria, so para exibicao
            "pc": rotulo(v("ponto")),
            "mun": rotulo(v("municipio")),
            "esc": norm_escritorio(v("escritorio")),
            "rt": norm_tecnico(v("tecnico")),
            "prod": titulo(texto(v("produtor"))) if texto(v("produtor")).isupper() else texto(v("produtor")),
            "pid": pid,
            "sexo": rotulo(v("sexo")),
            "ec": norm_estado_civil(v("civil")),
            # "Não" aqui significa "nao possui DAP" — nao pode ser tratado como ausencia
            "dap": texto(v("dap")) if texto(v("dap")) in ("Sim", "Não", "Vencida") else NAO_INF,
            "assoc": rotulo(v("assoc")),
            "loc": texto(v("endereco")),
            "propr": rotulo(v("propriedade")),
            "cult": culturas,
            "ha": round(numero(v("total_mec")), 2),
            "hrs": round(numero(v("horas")), 2),
            "ac": int(numero(v("acudes"))),
            "tt": rotulo(v("tipo_trator")),
            "maq": extrair(v("trator"), MAQUINAS) or extrair(v("maquina"), MAQUINAS),
            "impl": extrair(v("tipo_impl"), IMPLEMENTOS) or extrair(v("nome_impl"), IMPLEMENTOS),
            "dae": dae,
            "form": form,
            "obs": texto(v("obs")),
        })

    registros.sort(key=lambda x: x["d"])

    meta = {
        "arquivo": os.path.basename(XLSX),
        "aba": ABA,
        "gerado_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "registros": len(registros),
        "produtores": len(cpfs),
        "periodo": [registros[0]["d"], registros[-1]["d"]] if registros else ["", ""],
        "qualidade": qualidade,
    }

    corpo = json.dumps({"meta": meta, "registros": registros},
                       ensure_ascii=False, separators=(",", ":"))
    with io.open(SAIDA, "w", encoding="utf-8", newline="\n") as f:
        f.write("/* Dados de mecanizacao e acudagem — SEAGRI/AC\n")
        f.write("   GERADO AUTOMATICAMENTE por tools/gerar_dados_mecanizacao.py\n")
        f.write("   Fonte: %s (aba \"%s\") — %s\n" % (meta["arquivo"], ABA, meta["gerado_em"]))
        f.write("   Nao editar a mao: rode o script novamente apos atualizar a planilha. */\n")
        f.write("window.DADOS_MECANIZACAO = ")
        f.write(corpo)
        f.write(";\n")

    print("OK -> %s" % SAIDA)
    print("  registros: %d | produtores distintos: %d | periodo: %s a %s"
          % (meta["registros"], meta["produtores"], meta["periodo"][0], meta["periodo"][1]))
    print("  qualidade: %s" % qualidade)
    print("  tamanho: %.1f KB" % (os.path.getsize(SAIDA) / 1024.0))


if __name__ == "__main__":
    main()
